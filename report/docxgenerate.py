#!/usr/bin/env python
# python 3
#pylint: disable=
##    @file:    docxgenerate.py
#     @name:    Luke Gary
#  @company:    [company]
#     @date:    2025/10/2
################################################################################
# @copyright
#   Copyright 2025 [company] as an  unpublished work.
#   All Rights Reserved.
#
# @license The information contained herein is confidential
#   property of [company]. The user, copying, transfer or
#   disclosure of such information is prohibited except
#   by express written agreement with [company].
################################################################################

"""
{ high-level module description }
"""

import argparse
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import pprint as pp


def get_args():
    parser = argparse.ArgumentParser(description='create certificate of analysis')
    parser.add_argument('--mac', help='mac address', default='MACMACMACMAC')
    parser.add_argument('--serial', help='serial number', default='SNSNSNSNSNSN')
    parser.add_argument('--dom', help='date of manufacture', default='DD/MM/YYYY')
    parser.add_argument('--opt-cal', default='')
    parser.add_argument('--opt-qc', default='')
    parser.add_argument('--therm-cal', default='')
    parser.add_argument('--therm-qc', default='')
    parser.add_argument('--mech-qc', default='')
    parser.add_argument('--bio-qc', default='')
    return parser.parse_args()


class CertificateOfAnalysisDocument:
    '''Word Document Analysis Cert
    '''
    document_revision = '0.3'
    document_width = 6.5
    h_line_width = 78

     def __init__(self, **kwargs):
        """
        init
        """
        pp.pprint(kwargs)
        self._doc = Document()
        self.font = self._doc.styles['Normal'].font
        self.font.name = kwargs.get('font', 'Calibri')

        self.make_header()
        self.make_footer()
        self.make_title()
        self.make_device_info(**kwargs)
        self.make_performance_results(**kwargs)
        self.make_qc_signature()
        if kwargs.get('ivd_device', None):
            self.make_ivd_disclaimer()
        else:
            self.make_ruo_disclaimer()

    def make_hline(self):
        """
        Makes a hline.
        """
        hline = ''
        for _ in range(self.h_line_width):
            hline += f'_'

        self._doc.add_paragraph(hline)

    def make_header(self):
        """
        Makes a header.
        """
        header_section = self._doc.sections[0]
        header = header_section.header

        header_table = header.add_table(1, 2, Inches(self.document_width))
        header_tab_cells = header_table.rows[0].cells
        header_tab_0 = header_tab_cells[0].add_paragraph('')

        k_header = header_tab_0.add_run()
        k_header.add_picture(
            'company-logo.png',
            width=Inches(2.5)
        )
        header_tab_1 = header_tab_cells[1].add_paragraph(
            '<company>.\n'+\
            '<company addr 1>\n'+\
            '<company addr 2>'\
        )
        header_tab_1.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def make_footer(self):
        """
        Makes a footer.
        """
        footer_section = self._doc.sections[0]
        footer = footer_section.footer

        footer_table = footer.add_table(1, 3, Inches(self.document_width))
        footer_tab_cells = footer_table.rows[0].cells
        footer_tab_cells[0].add_paragraph('www.company.com')
        footer_tab_cells[1].add_paragraph('© company. 2020')
        footer_tab_cells[2].add_paragraph(
            f'<product>® COA Rev {self.document_revision}'
        )
        footer_tab_cells[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_tab_cells[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_tab_cells[2].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def make_title(self):
        """
        Makes a title.
        """
        title_para = self._doc.add_paragraph()
        run_1 = title_para.add_run('Certificate of Analysis - ')
        run_1.bold = True
        run_1.font.size = Pt(14)
        run_2 = title_para.add_run('<product>® Device')
        run_2.bold = False
        run_2.font.size = Pt(14)
        self.make_hline()

    def make_device_info(self, **kwargs):
        """
        Makes a device information.
        """

        dev_para = self._doc.add_paragraph()
        h_mac_run = dev_para.add_run('Device MAC Address   : ')
        mac_run = dev_para.add_run(f"{kwargs.get('mac')}\n")
        h_serial_run = dev_para.add_run('Device Serial Number : ')
        serial_run = dev_para.add_run(f"{kwargs.get('serial')}\n")
        h_dom_run = dev_para.add_run('Date of Manufacture  : ')
        dom_run = dev_para.add_run(f"{kwargs.get('dom')}")

        h_serial_run.bold = False
        h_serial_run.font.size = Pt(11)
        h_mac_run.bold = False
        h_mac_run.font.size = Pt(11)
        h_dom_run.bold = False
        h_dom_run.font.size = Pt(11)
        mac_run.bold = False
        mac_run.font.size = Pt(11)
        serial_run.bold = False
        serial_run.font.size = Pt(11)
        dom_run.bold = False
        dom_run.font.size = Pt(11)

        self.make_hline()

    def make_performance_results(self, **kwargs):
        """
        Makes performance results.
        """
        perf_table = self._doc.add_table(rows=7, cols=2)
        perf_table.cell(0, 1).text = 'Performance Results'
        perf_table.cell(1, 0).text = 'Optical Calibration'
        perf_table.cell(2, 0).text = 'Optical QC'
        perf_table.cell(3, 0).text = 'Thermal Calibration'
        perf_table.cell(4, 0).text = 'Thermal QC'
        perf_table.cell(5, 0).text = 'Mechanical QC'
        perf_table.cell(6, 0).text = 'qPCR QC'

        perf_table.cell(1, 1).text = kwargs.get('opt_cal', 'ERROR')
        perf_table.cell(2, 1).text = kwargs.get('opt_qc', 'ERROR')
        perf_table.cell(3, 1).text = kwargs.get('therm_cal', 'ERROR')
        perf_table.cell(4, 1).text = kwargs.get('therm_qc', 'ERROR')
        perf_table.cell(5, 1).text = kwargs.get('mech_qc', 'ERROR')
        perf_table.cell(6, 1).text = kwargs.get('bio_qc', 'ERROR')

        foot_note = self._doc.add_paragraph()
        foot_note_run = foot_note.add_run(
            'This unit has been calibrated and has passed all tests per the Manufacturing '+\
            'Specifications set forth by [company]'
        )
        foot_note_run.bold = False
        foot_note_run.font.size = Pt(9)

    def make_qc_signature(self):
        """
        Makes a qc signature.
        """
        sig_section = self._doc.add_table(1, 2)
        sig_col = sig_section.rows[0].cells
        sig_element = sig_col[1]

        sig_line = sig_element.add_paragraph()
        sig_line.add_run('_____________________________________\n')
        sig_line.add_run(' Quality Control / Quality Assurance')
        sig_line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def make_ivd_disclaimer(self):
        """
        Makes an ivd disclaimer.
        """
        pass

    def make_ruo_disclaimer(self):
        """
        Makes a ruo disclaimer.
        """
        disclaimer_par = self._doc.add_paragraph()
        disclaimer = disclaimer_par.add_run()
        disclaimer.text = \
            '\nThis product was tested according to {company}. internal procedures,'+\
            ' and the results showed that the product performed within specifications.\n'+\
            'The parts and supplies used to manufacture this product conform in all '+\
            'respects to {company}. product requirements, including specification, '+\
            'drawings, preservation, packaging, packing, marketing requirements and '+\
            'physical item identification (part number).\n'+\
            'This product was found to meet its functional and performance '+\
            'specifications set forth in {company}. internal procedures.  '+\
            'Biomeme, Inc. maintains possession of test and assembly reports as part '+\
            'of {company}’s Quality System.\n'+\
            'For Research Use Only. Not for use in Diagnostic Procedures. '.format(company=self.company)

        disclaimer.font.size = Pt(9)

    def save(self, *args, **kwargs):
        """
        save document
        :param      args:    The arguments
        :type       args:    list
        :param      kwargs:  The keywords arguments
        :type       kwargs:  dictionary
        """
        self._doc.save(*args, **kwargs)


def main():

    args = get_args()

    document = CertificateOfAnalysisDocument(**args.__dict__)
    document.save(f'test-{datetime.now().strftime("%m-%d-%Y_%H-%M-%S")}.docx')


if __name__ == '__main__':
    main()